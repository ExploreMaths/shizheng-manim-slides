# -*- coding: utf-8 -*-
"""把 时政演讲稿_动画标注版.md 转成排版良好的 docx（手持演讲稿 · 动画标注版）。"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:/Users/kylez/Doubao/chats/2026-10-17/new-chat/shizheng_manim/时政演讲稿_动画标注版.md"
OUT = r"C:/Users/kylez/Doubao/chats/2026-10-17/new-chat/shizheng_manim/时政演讲稿_动画标注版.docx"

RED = RGBColor(0xC0, 0x3A, 0x2B)
BLUE = RGBColor(0x1F, 0x6F, 0xB5)   # 动画标注用蓝色
INK = RGBColor(0x21, 0x1D, 0x18)
GRAY = RGBColor(0x7A, 0x71, 0x66)


def set_run(run, size=12, bold=False, color=None, ea="宋体", latin="Times New Roman", italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ea)


def add_runs(para, text, size=12, base_bold=False, ea="宋体"):
    """解析 **bold** 行内标记，并把 【动画▶…】 整段染成蓝色。"""
    # 先按 **bold** 切分
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if seg == "":
            continue
        is_bold = base_bold or (i % 2 == 1)
        # 再在每个 seg 内找 【动画▶…】 片段
        parts = re.split(r"(【动画▶[^】]*】)", seg)
        for p in parts:
            if p == "":
                continue
            r = para.add_run(p)
            if p.startswith("【动画▶"):
                set_run(r, size=size, bold=True, color=BLUE, ea=ea)
            else:
                set_run(r, size=size, bold=is_bold, color=None, ea=ea)


doc = Document()

sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.4)
sec.right_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
normal.paragraph_format.line_spacing = 1.4

# 页脚页码
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer_p.add_run("第 ")
set_run(r, size=9, color=GRAY)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer_p._p.append(fld)
r = footer_p.add_run(" 页")
set_run(r, size=9, color=GRAY)

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

i = 0
while i < len(lines):
    line = lines[i].rstrip()

    if line.strip() == "" or line.strip() == "---":
        i += 1
        continue

    # 表格
    if line.lstrip().startswith("|"):
        rows = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):
                rows.append(cells)
            i += 1
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            for ci, cell_text in enumerate(row):
                cell = table.cell(ri, ci)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                add_runs(cell.paragraphs[0], cell_text, size=10.5)
                if ri == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "F1E9DC")
                    cell._tc.get_or_add_tcPr().append(shd)
        doc.add_paragraph()
        continue

    # 一级标题
    if line.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(line[2:].strip())
        set_run(r, size=17, bold=True, color=INK, ea="黑体")
        i += 1
        continue

    # 二级标题 = 页码/章节
    if line.startswith("## "):
        text = line[3:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        set_run(r, size=13.5, bold=True, color=RED if "翻页" in text else INK, ea="黑体")
        i += 1
        continue

    # 整行加粗的说明行
    m = re.fullmatch(r"\*\*(.+?)\*\*", line.strip())
    if m:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, line.strip(), size=11)
        for r0 in p.runs:
            set_run(r0, size=11, bold=r0.font.bold, color=INK, ea="宋体")
        i += 1
        continue

    # 普通段落（含行内加粗 + 动画标注）
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_after = Pt(4)
    add_runs(p, line.strip(), size=12)
    i += 1

doc.save(OUT)
print("saved:", OUT)
