# -*- coding: utf-8 -*-
"""
把 assets/slides/ 下的分步 PNG 打包成一个 PPTX。
每张 PNG 占满一页 16:9 幻灯片，翻页笔每按一次 = 一个动画步。
输出：shizheng_manim_slides.pptx
"""
import os
import re
from pptx import Presentation
from pptx.util import Inches, Emu

SLIDES_DIR = r"C:\Users\kylez\Doubao\chats\2026-10-17\new-chat\shizheng_manim\assets\slides"
OUT = r"C:\Users\kylez\Doubao\chats\2026-10-17\new-chat\shizheng_manim\shizheng_manim_slides.pptx"

# 收集所有 PNG 并按 s{页码}_step{步} 排序
files = [f for f in os.listdir(SLIDES_DIR) if f.endswith(".png")]
def keyfn(f):
    m = re.match(r"s(\d+)_step(\d+)", f)
    return (int(m.group(1)), int(m.group(2)))
files.sort(key=keyfn)

print(f"共 {len(files)} 张分步图：")
for f in files:
    print("  ", f)

# 创建 16:9 PPTX（13.333 x 7.5 英寸）
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]  # 空白版式

for f in files:
    slide = prs.slides.add_slide(blank_layout)
    img_path = os.path.join(SLIDES_DIR, f)
    # 图片铺满整页
    slide.shapes.add_picture(
        img_path,
        left=0, top=0,
        width=prs.slide_width, height=prs.slide_height,
    )

prs.save(OUT)
print(f"\n已保存: {OUT}")
print(f"共 {len(prs.slides)} 页（= {len(files)} 个分步状态）")
